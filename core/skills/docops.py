"""Document helpers (plan §34 "Document read", §56 Document capability; Phase 20):
read PDF / Word / Excel / CSV / text, write Word / Markdown / text reports and
Excel sheets, merge / split PDFs. Pure Python (pypdf, python-docx, openpyxl).

Document text is UNTRUSTED DATA (plan §19) — it is returned as data, never acted on."""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

READABLE = (".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".txt", ".md", ".json", ".log")
MAX_READ_CHARS = 60_000
MAX_SHEET_ROWS = 2000


def read_document(path: Path, max_chars: int = MAX_READ_CHARS) -> tuple[str, dict[str, Any]]:
    """(text, meta) — meta has pages / paragraphs / sheets."""
    ext = path.suffix.lower()
    meta: dict[str, Any] = {"type": ext.lstrip(".")}
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        meta["pages"] = len(reader.pages)
        chunks, used = [], 0
        for i, page in enumerate(reader.pages, start=1):
            t = (page.extract_text() or "").strip()
            if t:
                chunks.append(f"--- page {i} ---\n{t}")
                used += len(t)
            if used >= max_chars:
                break
        text = "\n".join(chunks)
        if not text.strip():
            meta["note"] = "no text layer (a scanned PDF needs OCR)"
    elif ext == ".docx":
        import docx
        d = docx.Document(str(path))
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        meta["paragraphs"] = len(paras)
        rows = []
        for table in d.tables:
            for r in table.rows:
                rows.append(" | ".join(c.text.strip() for c in r.cells))
        meta["tables"] = len(d.tables)
        text = "\n".join(paras + (["", "Tables:", *rows] if rows else []))
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        meta["sheets"] = wb.sheetnames
        lines = []
        for ws in wb.worksheets:
            lines.append(f"--- sheet {ws.title} ---")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_SHEET_ROWS:
                    lines.append("… (more rows)")
                    break
                if any(v is not None for v in row):
                    lines.append(" | ".join("" if v is None else str(v) for v in row))
        wb.close()
        text = "\n".join(lines)
    elif ext == ".csv":
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as f:
            rows = []
            for i, r in enumerate(csv.reader(f)):
                if i >= MAX_SHEET_ROWS:
                    rows.append("… (more rows)")
                    break
                rows.append(" | ".join(r))
        meta["rows"] = len(rows)
        text = "\n".join(rows)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    meta["chars"] = len(text)
    meta["truncated"] = len(text) > max_chars
    return text[:max_chars], meta


_HEADING = re.compile(r"^(#{1,4})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
_CHECK = re.compile(r"^\s*[-*]\s*\[( |x|X)\]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")


def write_docx(path: Path, title: str | None, content: str) -> int:
    """Markdown-ish text → Word: # headings, - bullets, 1. numbers, - [ ] checklist,
    blank-line paragraphs. Returns the number of blocks written."""
    import docx
    d = docx.Document()
    if title:
        d.add_heading(title, level=0)
    blocks = 0
    for line in content.splitlines():
        if not line.strip():
            continue
        if (m := _HEADING.match(line)) is not None:
            d.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
        elif (m := _CHECK.match(line)) is not None:
            d.add_paragraph(("☑ " if m.group(1).lower() == "x" else "☐ ") + m.group(2))
        elif (m := _BULLET.match(line)) is not None:
            d.add_paragraph(m.group(1), style="List Bullet")
        elif (m := _NUMBERED.match(line)) is not None:
            d.add_paragraph(m.group(1), style="List Number")
        else:
            d.add_paragraph(line.strip())
        blocks += 1
    d.save(str(path))
    return blocks


def write_sheet(path: Path, rows: list[list[Any]], sheet: str = "Sheet1") -> int:
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet[:31]
    for r in rows:
        ws.append(r)
    if rows:
        for cell in ws[1]:
            cell.font = Font(bold=True)            # first row = header
    wb.save(str(path))
    return len(rows)


def parse_pages(spec: str, total: int) -> list[int]:
    """'1-3,5' → [0, 1, 2, 4] (0-based), within 1..total."""
    pages: list[int] = []
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            start, end = int(a), int(b or total)
        else:
            start = end = int(part)
        if start < 1 or end > total or start > end:
            raise ValueError(f"pages {part} outside 1–{total}")
        pages.extend(range(start - 1, end))
    if not pages:
        raise ValueError("no pages given")
    return pages
